#!/usr/bin/env python3
"""
Document Parser for EDCAT Vera
Extracts text and quality-related entities from PDF, DOCX, and spreadsheet files

Usage:
    python parse-documents.py --input-dir ./data/raw_documents --output-dir ./data/parsed

Requirements:
    pip install PyPDF2 python-docx openpyxl pandas
"""

import os
import sys
import json
import argparse
import re
from pathlib import Path
from typing import List, Dict, Optional
from datetime import datetime

try:
    import PyPDF2
except ImportError:
    print("❌ Missing PyPDF2. Install with: pip install PyPDF2")
    sys.exit(1)

try:
    from docx import Document
except ImportError:
    print("❌ Missing python-docx. Install with: pip install python-docx")
    sys.exit(1)

try:
    import pandas as pd
except ImportError:
    print("❌ Missing pandas. Install with: pip install pandas openpyxl")
    sys.exit(1)


class DocumentParser:
    """Parse documents and extract quality-related entities"""
    
    def __init__(self, input_dir: str, output_dir: str):
        self.input_dir = Path(input_dir)
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)
        
        # Statistics
        self.stats = {
            'total_files': 0,
            'parsed': 0,
            'errors': 0,
            'total_chars': 0,
            'total_clauses': 0,
            'total_standards': 0,
            'total_procedures': 0,
            'total_definitions': 0
        }
    
    def parse_all_documents(self) -> List[Dict]:
        """Parse all documents and extract text"""
        documents = []
        
        if not self.input_dir.exists():
            print(f"❌ Input directory not found: {self.input_dir}")
            return documents
        
        # Get all supported files
        supported_extensions = {'.pdf', '.docx', '.xlsx', '.xls', '.csv'}
        files = [
            f for f in self.input_dir.iterdir()
            if f.is_file() and f.suffix.lower() in supported_extensions
        ]
        
        self.stats['total_files'] = len(files)
        print(f"📁 Found {len(files)} documents to parse\n")
        
        for idx, filepath in enumerate(files, 1):
            filename = filepath.name
            print(f"[{idx}/{len(files)}] Processing: {filename}")
            
            try:
                # Parse based on file type
                if filepath.suffix.lower() == '.pdf':
                    text = self.parse_pdf(filepath)
                elif filepath.suffix.lower() == '.docx':
                    text = self.parse_docx(filepath)
                elif filepath.suffix.lower() in {'.xlsx', '.xls', '.csv'}:
                    text = self.parse_spreadsheet(filepath)
                else:
                    continue
                
                if not text or len(text.strip()) < 10:
                    print(f"   ⚠️  Skipped: Empty or too short")
                    continue
                
                # Extract quality-related entities
                clauses = self.extract_clauses(text)
                standards = self.extract_standards(text)
                procedures = self.extract_procedures(text)
                definitions = self.extract_definitions(text)
                
                doc = {
                    'filename': filename,
                    'filepath': str(filepath),
                    'text': text,
                    'clauses': clauses,
                    'standards': standards,
                    'procedures': procedures,
                    'definitions': definitions,
                    'length': len(text),
                    'parsed_at': datetime.now().isoformat()
                }
                
                documents.append(doc)
                self.stats['parsed'] += 1
                self.stats['total_chars'] += len(text)
                self.stats['total_clauses'] += len(clauses)
                self.stats['total_standards'] += len(standards)
                self.stats['total_procedures'] += len(procedures)
                self.stats['total_definitions'] += len(definitions)
                
                print(f"   ✅ Parsed: {len(text):,} chars, {len(clauses)} clauses, {len(standards)} standards")
                
            except Exception as e:
                print(f"   ❌ Error: {e}")
                self.stats['errors'] += 1
        
        # Save parsed documents
        output_file = self.output_dir / 'parsed_documents.json'
        with open(output_file, 'w', encoding='utf-8') as f:
            json.dump({
                'metadata': {
                    'parsed_at': datetime.now().isoformat(),
                    'input_dir': str(self.input_dir),
                    'stats': self.stats
                },
                'documents': documents
            }, f, indent=2, ensure_ascii=False)
        
        self.print_summary()
        
        return documents
    
    def parse_pdf(self, filepath: Path) -> str:
        """Extract text from PDF"""
        text = ""
        try:
            with open(filepath, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                for page_num, page in enumerate(reader.pages, 1):
                    try:
                        page_text = page.extract_text()
                        if page_text:
                            text += page_text + "\n"
                    except Exception as e:
                        print(f"      ⚠️  Error extracting page {page_num}: {e}")
        except Exception as e:
            raise Exception(f"PDF parsing failed: {e}")
        return text.strip()
    
    def parse_docx(self, filepath: Path) -> str:
        """Extract text from DOCX"""
        try:
            doc = Document(filepath)
            paragraphs = []
            for para in doc.paragraphs:
                if para.text.strip():
                    paragraphs.append(para.text)
            
            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            
            return "\n".join(paragraphs)
        except Exception as e:
            raise Exception(f"DOCX parsing failed: {e}")
    
    def parse_spreadsheet(self, filepath: Path) -> str:
        """Extract text from CSV/XLSX"""
        try:
            if filepath.suffix.lower() == '.csv':
                df = pd.read_csv(filepath, encoding='utf-8', errors='ignore')
            else:
                df = pd.read_excel(filepath, engine='openpyxl')
            
            # Convert to readable text format
            text_parts = []
            
            # Add column headers
            if not df.empty:
                text_parts.append("Columns: " + ", ".join(df.columns.astype(str)))
                text_parts.append("")
                
                # Add data rows (limit to first 100 rows for very large files)
                max_rows = min(100, len(df))
                for idx, row in df.head(max_rows).iterrows():
                    row_text = " | ".join(str(val) for val in row.values if pd.notna(val))
                    if row_text:
                        text_parts.append(f"Row {idx + 1}: {row_text}")
            
            return "\n".join(text_parts)
        except Exception as e:
            raise Exception(f"Spreadsheet parsing failed: {e}")
    
    def extract_clauses(self, text: str) -> List[str]:
        """Extract clause references (e.g., 8.5.1, 4.2.1, 21.137(n))"""
        clauses = []
        
        # ISO/AS style clauses: 8.5.1, 4.2.1, etc.
        pattern1 = r'\b(\d+\.\d+(?:\.\d+)*(?:\.\d+)?)\b'
        matches1 = re.findall(pattern1, text)
        clauses.extend(matches1)
        
        # Regulatory style: 21.137(n), Part 21.137, etc.
        pattern2 = r'(?:Part\s+)?(\d+\.\d+(?:\([a-z]\))?)'
        matches2 = re.findall(pattern2, text, re.IGNORECASE)
        clauses.extend(matches2)
        
        # Clause references with "Clause" prefix
        pattern3 = r'[Cc]lause\s+(\d+\.\d+(?:\.\d+)*)'
        matches3 = re.findall(pattern3, text)
        clauses.extend(matches3)
        
        # Remove duplicates and sort
        unique_clauses = sorted(list(set(clauses)), key=lambda x: (
            len(x.split('.')),  # Sort by depth first
            [int(part) if part.isdigit() else 0 for part in x.split('.')]  # Then numerically
        ))
        
        return unique_clauses[:50]  # Limit to top 50
    
    def extract_standards(self, text: str) -> List[str]:
        """Extract standard references (AS9100D, ISO 9001, etc.)"""
        standards = []
        
        patterns = [
            r'AS9100[A-D]?',
            r'AS9110[A-D]?',
            r'AS9120[A-D]?',
            r'ISO\s*\d+(?::\d+)?',
            r'IATF\s*\d+',
            r'FAA\s*(?:Part\s*)?\d+(?:\.\d+)?',
            r'CASA\s*(?:Part\s*)?\d+(?:\.\d+)?',
            r'EASA\s*(?:Part\s*)?\d+(?:\.\d+)?',
            r'UK\s*CAA',
            r'FDA\s*(?:Part\s*)?\d+(?:\.\d+)?',
            r'IEC\s*\d+',
            r'ISO\s*13485',
            r'ISO\s*14971',
            r'MDSAP',
            r'MDR',
            r'IVDR',
        ]
        
        for pattern in patterns:
            matches = re.findall(pattern, text, re.IGNORECASE)
            standards.extend(matches)
        
        # Normalize and deduplicate
        normalized = []
        seen = set()
        for std in standards:
            std_upper = std.upper()
            if std_upper not in seen:
                normalized.append(std)
                seen.add(std_upper)
        
        return sorted(normalized)
    
    def extract_procedures(self, text: str) -> List[str]:
        """Extract procedure names and references"""
        procedures = []
        
        # Look for procedure titles/headers
        keywords = [
            r'procedure[:\s]+([A-Z][^\n]{10,100})',
            r'work\s+instruction[:\s]+([A-Z][^\n]{10,100})',
            r'process[:\s]+([A-Z][^\n]{10,100})',
            r'method[:\s]+([A-Z][^\n]{10,100})',
            r'guideline[:\s]+([A-Z][^\n]{10,100})',
            r'standard\s+operating\s+procedure[:\s]+([A-Z][^\n]{10,100})',
            r'sop[:\s]+([A-Z][^\n]{10,100})',
        ]
        
        for pattern in keywords:
            matches = re.findall(pattern, text, re.IGNORECASE)
            procedures.extend(matches)
        
        # Clean and limit
        cleaned = [p.strip()[:100] for p in procedures if len(p.strip()) > 10]
        return list(set(cleaned))[:50]  # Limit to 50
    
    def extract_definitions(self, text: str) -> List[tuple]:
        """Extract definitions and explanations"""
        definitions = []
        
        # Pattern 1: "Term is/means/refers to/defined as..."
        pattern1 = r'([A-Z][A-Za-z]+(?:\s+[A-Z][A-Za-z]+)*)\s+(?:is|means|refers to|defined as|shall be)\s+([^.\n]{20,200})'
        matches1 = re.findall(pattern1, text)
        definitions.extend(matches1)
        
        # Pattern 2: "Term: definition"
        pattern2 = r'([A-Z][A-Za-z]+(?:\s+[A-Z][A-Za-z]+)*):\s+([^.\n]{20,200})'
        matches2 = re.findall(pattern2, text)
        definitions.extend(matches2)
        
        # Pattern 3: Acronym definitions: "APQP (Advanced Product Quality Planning)"
        pattern3 = r'([A-Z]{2,10})\s*\(([A-Z][^)]{10,100})\)'
        matches3 = re.findall(pattern3, text)
        definitions.extend(matches3)
        
        # Clean and limit
        cleaned = [(term.strip(), defn.strip()) for term, defn in definitions 
                  if len(term.strip()) > 2 and len(defn.strip()) > 10]
        
        return cleaned[:100]  # Limit to 100
    
    def print_summary(self):
        """Print parsing summary"""
        print("\n" + "="*60)
        print("📊 Parsing Summary")
        print("="*60)
        print(f"   📁 Total files: {self.stats['total_files']}")
        print(f"   ✅ Parsed: {self.stats['parsed']}")
        print(f"   ❌ Errors: {self.stats['errors']}")
        print(f"   📝 Total text: {self.stats['total_chars']:,} characters")
        print(f"   📋 Clauses found: {self.stats['total_clauses']}")
        print(f"   📚 Standards found: {self.stats['total_standards']}")
        print(f"   📄 Procedures found: {self.stats['total_procedures']}")
        print(f"   📖 Definitions found: {self.stats['total_definitions']}")
        print(f"   💾 Output: {self.output_dir / 'parsed_documents.json'}")
        print("="*60)


def main():
    parser = argparse.ArgumentParser(
        description='Parse documents and extract quality-related entities',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
  # Parse documents from default directories
  python parse-documents.py --input-dir ./data/raw_documents --output-dir ./data/parsed
  
  # Parse specific directory
  python parse-documents.py -i /path/to/documents -o /path/to/output
        """
    )
    
    parser.add_argument(
        '-i', '--input-dir',
        default='./data/raw_documents',
        help='Input directory with documents (default: ./data/raw_documents)'
    )
    
    parser.add_argument(
        '-o', '--output-dir',
        default='./data/parsed',
        help='Output directory for parsed JSON (default: ./data/parsed)'
    )
    
    args = parser.parse_args()
    
    print("="*60)
    print("Document Parser for EDCAT Vera")
    print("="*60)
    print()
    
    parser = DocumentParser(args.input_dir, args.output_dir)
    documents = parser.parse_all_documents()
    
    if documents:
        print(f"\n✅ Successfully parsed {len(documents)} documents!")
        print(f"📁 Output saved to: {args.output_dir}/parsed_documents.json")
        print("\n💡 Next step: Run create-training-data.py to generate training pairs")
    else:
        print("\n⚠️  No documents were parsed. Check input directory and file formats.")


if __name__ == "__main__":
    main()

