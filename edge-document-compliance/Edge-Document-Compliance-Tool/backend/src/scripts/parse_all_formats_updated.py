#!/usr/bin/env python3
import os
import json
from pathlib import Path

quality_sources = [
    "D:\\",
    r"C:\Users\oleab\OneDrive\Conform-Edge\EDCAT-Administrator\data\quality_docs_all",
    r"C:\Users\oleab\OneDrive\Conform-Edge\EDCAT-Administrator\backend\EDCAT AIP\Requirements",
]

output_dir = r"D:\quality-llm\data"
parsed_file = os.path.join(output_dir, "parsed_all_comprehensive.json")

os.makedirs(output_dir, exist_ok=True)

print("\nPARSING ALL QUALITY DOCUMENTS (WITH HTML)\n")

parsed_docs = []
parsed_count = 0

target_extensions = {'.pdf', '.docx', '.doc', '.xlsx', '.xls', '.csv', '.txt', '.md', '.pptx', '.ppt', '.rtf', '.dotx', '.xlsm', '.htm', '.html'}

for source in quality_sources:
    if not os.path.exists(source):
        continue
    
    print(f"Processing: {source}")
    
    for filepath in Path(source).rglob('*'):
        if not filepath.is_file():
            continue
        
        if filepath.stat().st_size == 0:
            continue
        
        ext = filepath.suffix.lower()
        
        if ext not in target_extensions:
            continue
        
        if any(skip in str(filepath).lower() for skip in ['node_modules', '.git', '__pycache__']):
            continue
        
        text = ""
        
        try:
            if ext in ['.txt', '.md', '.rtf', '.htm', '.html']:
                with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                    text = f.read()
            
            elif ext == '.pdf':
                import PyPDF2
                with open(filepath, 'rb') as f:
                    reader = PyPDF2.PdfReader(f)
                    for page in reader.pages:
                        text += page.extract_text() + "\n"
            
            elif ext in ['.docx', '.dotx', '.doc']:
                from docx import Document
                doc = Document(filepath)
                text = '\n'.join([p.text for p in doc.paragraphs])
            
            elif ext in ['.xlsx', '.xls', '.xlsm']:
                import pandas as pd
                df = pd.read_excel(filepath)
                text = df.to_string()
            
            elif ext == '.csv':
                import pandas as pd
                df = pd.read_csv(filepath)
                text = df.to_string()
            
            elif ext in ['.pptx', '.ppt']:
                from pptx import Presentation
                prs = Presentation(filepath)
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if hasattr(shape, "text"):
                            text += shape.text + "\n"
            
            if len(text) > 50:
                parsed_docs.append({
                    'filename': filepath.name,
                    'filepath': str(filepath),
                    'source': source,
                    'text': text[:30000],
                    'size': len(text)
                })
                parsed_count += 1
                
                if parsed_count % 500 == 0:
                    print(f"   Parsed {parsed_count} documents...")
        
        except Exception as e:
            pass

print(f"\nRESULTS:")
print(f"   Total: {parsed_count} documents\n")

with open(parsed_file, 'w') as f:
    json.dump(parsed_docs, f, indent=2)

print(f"Saved: {parsed_file}\n")